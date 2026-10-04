"""Stage the local CRM workbooks without sending them anywhere.

The output is intentionally ignored by git because it contains private prospect data.
Each normalized record keeps its exact source row in ``raw`` plus provenance fields,
so a later authenticated importer can create prospects, contacts, branches and tags
without losing the original Excel values.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "staging"

PIPELINE_RECEIVED = "clientes recibidos"
DEFAULT_COUNTRY = "México"


def clean_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    text = str(value).replace("\x00", "").strip()
    return text or None


def key_text(value: Any) -> str:
    text = clean_value(value) or ""
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.replace("�", " ")
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def slug(value: Any) -> str:
    return re.sub(r"\s+", "-", key_text(value)).strip("-")


def contains_mojibake(value: Any) -> bool:
    text = clean_value(value) or ""
    return "�" in text or "Ã" in text or "Â" in text


def headers_for(row: Iterable[Any]) -> list[str]:
    result: list[str] = []
    for index, value in enumerate(row):
        result.append(clean_value(value) or f"column_{index + 1}")
    return result


def row_dict(headers: list[str], values: Iterable[Any]) -> dict[str, Any]:
    cells = list(values)
    return {headers[i]: clean_value(cells[i]) if i < len(cells) else None for i in range(len(headers))}


def find_header(record: dict[str, Any], *names: str) -> Any:
    wanted = {key_text(name) for name in names}
    for key, value in record.items():
        if key_text(key) in wanted:
            return value
    return None


def first_nonempty(*values: Any) -> Any:
    for value in values:
        if clean_value(value):
            return clean_value(value)
    return None


def clean_company_name(value: Any) -> str | None:
    text = clean_value(value)
    if not text:
        return None
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if lines and key_text(lines[0]) == "nombre de la empresa":
        lines = lines[1:]
    return "\n".join(lines) or None


def is_placeholder_company(value: Any) -> bool:
    normalized = key_text(value)
    return bool(re.fullmatch(r"prospecto yucatan \d+", normalized) or re.fullmatch(r"comercial \d+", normalized) or normalized in {"nombre", "nombre de la empresa", "razon social"})


def extract_emails(text: Any) -> list[str]:
    return sorted(set(re.findall(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", clean_value(text) or "", re.I)))


def extract_urls(text: Any) -> list[str]:
    values = re.findall(r"https?://[^\s,;]+|(?<!@)(?:www\.)[a-z0-9.-]+\.[a-z]{2,}(?:/[^\s,;]*)?", clean_value(text) or "", re.I)
    return sorted(set(value.rstrip(".)]") for value in values))


def extract_phones(text: Any) -> list[str]:
    candidates = re.findall(r"(?:\+?\d[\d .()/-]{7,}\d)(?:\s*(?:ext\.?|x)\s*\d+)?", clean_value(text) or "", re.I)
    result: list[str] = []
    for candidate in candidates:
        normalized = re.sub(r"\s+", " ", candidate).strip(" ,;.")
        digits = re.sub(r"\D", "", normalized)
        if len(digits) >= 8:
            result.append(normalized)
    return sorted(set(result))


def split_location(value: Any) -> dict[str, Any]:
    text = clean_value(value)
    if not text:
        return {"address": None, "city": None, "state": None, "postal_code": None}
    cp = re.search(r"\b(\d{5})\b", text)
    states = ["Yucatán", "Quintana Roo", "Campeche", "Puebla", "Jalisco", "Guanajuato", "Veracruz"]
    state = next((item for item in states if key_text(item) in key_text(text)), None)
    city = None
    if "," in text:
        chunks = [chunk.strip() for chunk in text.split(",") if chunk.strip()]
        if len(chunks) >= 2:
            city = chunks[-2] if state and len(chunks) >= 3 else chunks[-1]
    return {"address": text, "city": city, "state": state, "postal_code": cp.group(1) if cp else None}


def priority(value: Any) -> str:
    current = key_text(value)
    if current == "alta":
        return "alta"
    if current == "baja":
        return "baja"
    return "media"


def suggested_stage(raw_status: Any, notes: Any, conversation: Any) -> tuple[str, str]:
    raw = key_text(raw_status)
    text = key_text(f"{clean_value(notes) or ''} {clean_value(conversation) or ''}")
    if any(token in raw for token in ("sin interes", "no interesado", "no le interesa")):
        return "clientes no interesados", "status_raw"
    if any(token in raw for token in ("contactado", "visitado", "atacado")):
        return "clientes contactados", "status_raw"
    if "esper" in raw or any(token in text for token in ("esperando", "en espera", "pendiente de")):
        return "clientes en espera", "status_or_notes"
    if "interes" in raw or any(token in text for token in ("interesado", "interesada", "le interesa")):
        return "clientes interesados", "status_or_notes"
    if "prospect" in raw or "seguimiento" in raw:
        return "clientes recibidos", "status_raw"
    return PIPELINE_RECEIVED, "default_needs_review"


def normalized_record(*, source_file: str, sheet: str, row_number: int, raw: dict[str, Any], company: Any,
                      trade_name: Any = None, contact: Any = None, email: Any = None, phone: Any = None,
                      website: Any = None, location: Any = None, industry: Any = None, raw_status: Any = None,
                      raw_priority: Any = None, follow_up: Any = None, notes: Any = None,
                      conversation: Any = None, rfc: Any = None, extra: dict[str, Any] | None = None) -> dict[str, Any] | None:
    company_value = clean_company_name(first_nonempty(company, trade_name))
    if not company_value or is_placeholder_company(company_value):
        return None
    contact_raw = first_nonempty(contact)
    emails = extract_emails(first_nonempty(email, contact_raw))
    urls = extract_urls(first_nonempty(website, contact_raw))
    phones = extract_phones(first_nonempty(phone, contact_raw))
    location_parts = split_location(location)
    suggested, suggestion_reason = suggested_stage(raw_status, notes, conversation)
    flags = []
    if any(contains_mojibake(value) for value in raw.values()):
        flags.append("encoding_mojibake")
    if not emails:
        flags.append("missing_email")
    if not phones:
        flags.append("missing_phone")
    if not location_parts["address"]:
        flags.append("missing_location")
    normalized_name = key_text(company_value)
    record_key = f"{source_file}:{sheet}:{row_number}"
    return {
        "source_file": source_file,
        "source_sheet": sheet,
        "source_row": row_number,
        "record_key": record_key,
        "company_name": company_value,
        "trade_name": clean_value(trade_name),
        "contact_name": clean_value(contact) if contact and "\n" not in str(contact) else None,
        "contact_raw": contact_raw,
        "emails": emails,
        "phones": phones,
        "websites": sorted(set(urls)),
        "address": location_parts["address"],
        "city": location_parts["city"],
        "state": location_parts["state"],
        "postal_code": location_parts["postal_code"],
        "country": DEFAULT_COUNTRY,
        "industry": first_nonempty(industry, "Por clasificar"),
        "priority": priority(raw_priority),
        "stage": PIPELINE_RECEIVED,
        "suggested_stage": suggested,
        "stage_suggestion_reason": suggestion_reason,
        "status_raw": clean_value(raw_status),
        "next_follow_up_raw": clean_value(follow_up),
        "notes": clean_value(notes),
        "conversation_notes": clean_value(conversation),
        "rfc": clean_value(rfc),
        "dedupe_key": f"{normalized_name}|{key_text(location_parts['city'] or location_parts['state'] or '')}",
        "quality_flags": flags,
        "extra": extra or {},
        "raw": raw,
    }


def append_record(records: list[dict[str, Any]], **kwargs: Any) -> None:
    record = normalized_record(**kwargs)
    if record:
        records.append(record)


def read_clientes(path: Path, records: list[dict[str, Any]], stats: dict[str, Any]) -> None:
    wb = load_workbook(path, read_only=True, data_only=True)
    skip = {"an�lisis comercial", "notas"}
    for sheet in wb.sheetnames:
        if key_text(sheet) in {key_text(item) for item in skip}:
            stats["skipped_sheets"].append({"file": path.name, "sheet": sheet, "reason": "dashboard_or_notes"})
            continue
        ws = wb[sheet]
        rows = ws.iter_rows(min_row=3, values_only=True)
        headers = headers_for(next(rows))
        for row_number, values in enumerate(rows, start=4):
            raw = row_dict(headers, values)
            company = find_header(raw, "Compañía", "Razón social")
            append_record(records, source_file=path.name, sheet=sheet, row_number=row_number, raw=raw,
                          company=company, trade_name=find_header(raw, "Nombre", "Nombre Comercial"),
                          contact=find_header(raw, "Información de contacto"), location=find_header(raw, "Ubicación"),
                          industry=sheet, raw_status=find_header(raw, "Completado"),
                          follow_up=find_header(raw, "Fecha de seguimiento"), notes=find_header(raw, "Notas"),
                          conversation=find_header(raw, "Notas de la conversación"))


def read_merida(path: Path, records: list[dict[str, Any]], stats: dict[str, Any]) -> None:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Prospectos"]
    iterator = ws.iter_rows(values_only=True)
    headers = headers_for(next(iterator))
    for row_number, values in enumerate(iterator, start=2):
        raw = row_dict(headers, values)
        company = find_header(raw, "Empresa")
        if is_placeholder_company(company):
            stats["placeholder_rows_skipped"] += 1
            continue
        append_record(records, source_file=path.name, sheet=ws.title, row_number=row_number, raw=raw,
                      company=company, trade_name=find_header(raw, "Nombre Comercial"),
                      phone=find_header(raw, "Telefono"), website=find_header(raw, "Web"),
                      location=find_header(raw, "Direccion"), industry=find_header(raw, "Giro/Importa-Exporta"),
                      raw_priority=find_header(raw, "Prioridad"), rfc=find_header(raw, "RFC"),
                      extra={"ranking": find_header(raw, "Ranking")})


def read_immex(path: Path, records: list[dict[str, Any]], stats: dict[str, Any]) -> None:
    wb = load_workbook(path, read_only=True, data_only=True)
    for sheet in ("PROGRAMAS IMMEX", "PLANTAS IMMEX"):
        ws = wb[sheet]
        iterator = ws.iter_rows(min_row=5, values_only=True)
        headers = headers_for(next(iterator))
        for row_number, values in enumerate(iterator, start=6):
            raw = row_dict(headers, values)
            company = find_header(raw, "RAZÓN_SOCIAL")
            address = " ".join(filter(None, [find_header(raw, "CALLE"), find_header(raw, "NÚMERO_EXTERIOR"), find_header(raw, "COLONIA")]))
            location = ", ".join(filter(None, [address, find_header(raw, "CIUDAD"), find_header(raw, "MUNICIPIO"), find_header(raw, "ESTADO"), find_header(raw, "CP")]))
            append_record(records, source_file=path.name, sheet=sheet, row_number=row_number, raw=raw,
                          company=company, email=find_header(raw, "CORREO_ELECTRONICO"),
                          phone=find_header(raw, "TELEFONO", "TELÉFONO"), location=location,
                          industry=find_header(raw, "ACTIVIDAD", "TIPO_PROGRAMA"),
                          raw_status=find_header(raw, "ESTATUS_PROGRAMA"),
                          extra={"immex_number": find_header(raw, "NÚMERO_IMMEX"), "immex_year": find_header(raw, "AÑO_IMMEX"), "delegation": find_header(raw, "DELEGACIÓN")})


def read_hoja_trabajo(path: Path, records: list[dict[str, Any]], stats: dict[str, Any]) -> None:
    selected = {"Lu Prosp", "Prosp", "P.Puebla", "P.Sergio", "MRR 1", "MRR 2", "MRR 3"}
    wb = load_workbook(path, read_only=True, data_only=True)
    for sheet in wb.sheetnames:
        if sheet not in selected:
            stats["skipped_sheets"].append({"file": path.name, "sheet": sheet, "reason": "financial_or_operational_sheet"})
            continue
        ws = wb[sheet]
        all_rows = list(ws.iter_rows(values_only=True))
        header_index = next((i for i, values in enumerate(all_rows) if any(key_text(value) in {"razon social", "empresa o prospecto", "promotor monex"} for value in values if value)), None)
        if header_index is None:
            stats["skipped_sheets"].append({"file": path.name, "sheet": sheet, "reason": "no_prospect_header"})
            continue
        headers = headers_for(all_rows[header_index])
        for row_number, values in enumerate(all_rows[header_index + 1:], start=header_index + 2):
            raw = row_dict(headers, values)
            company = find_header(raw, "RAZON SOCIAL", "Empresa o prospecto")
            trade = find_header(raw, "NOMBRE COMERCIAL", "RELACION")
            contact = find_header(raw, "CONTACTO")
            email = find_header(raw, "CORREO")
            phone = find_header(raw, "TELEFONO", "TELEFONICO")
            website = find_header(raw, "PAGINA WEB", "SITIO WEB")
            location = find_header(raw, "DIRECCION", "Ciudad")
            notes = find_header(raw, "COMENTARIOS", "Nota")
            industry = find_header(raw, "Descripción", "GENERA 2022", "CL")
            append_record(records, source_file=path.name, sheet=sheet, row_number=row_number, raw=raw,
                          company=company, trade_name=trade, contact=contact, email=email, phone=phone,
                          website=website, location=location, industry=industry, notes=notes,
                          extra={"promotor": find_header(raw, "PROMOTORA:", "PROMOTOR/MONEX"), "source_campaign": find_header(raw, "PROSPECTO BUSCADO POR:")})


def checksum(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    records: list[dict[str, Any]] = []
    stats: dict[str, Any] = {"skipped_sheets": [], "placeholder_rows_skipped": 0}
    files = [
        (ROOT / "Clientes potenciales.xlsx", read_clientes),
        (ROOT / "Prospectos_Masari_Merida_220.xlsx", read_merida),
        (ROOT / "IMMEX_MAYO2023-DIRECTORIO_20230609-20230609.xlsx", read_immex),
        (ROOT / "Hoja de TRABAJO Masari.xlsx", read_hoja_trabajo),
    ]
    for path, reader in files:
        if not path.exists():
            raise FileNotFoundError(path)
        reader(path, records, stats)

    dedupe_groups: dict[str, list[str]] = defaultdict(list)
    for record in records:
        dedupe_groups[record["dedupe_key"]].append(record["record_key"])
    duplicate_groups = {key: values for key, values in dedupe_groups.items() if key.split("|", 1)[0] and len(values) > 1}

    coverage_fields = ["company_name", "trade_name", "contact_raw", "emails", "phones", "websites", "address", "notes", "conversation_notes"]
    coverage = {field: sum(bool(record.get(field)) for record in records) for field in coverage_fields}
    stats.update({
        "generated_at": datetime.now().astimezone().isoformat(),
        "source_files": [{"name": path.name, "bytes": path.stat().st_size, "sha256": checksum(path)} for path, _ in files],
        "records_total": len(records),
        "records_by_file": dict(Counter(record["source_file"] for record in records)),
        "records_by_sheet": dict(Counter(f"{record['source_file']}::{record['source_sheet']}" for record in records)),
        "records_by_suggested_stage": dict(Counter(record["suggested_stage"] for record in records)),
        "records_with_mojibake_flag": sum("encoding_mojibake" in record["quality_flags"] for record in records),
        "duplicate_groups": duplicate_groups,
        "duplicate_group_count": len(duplicate_groups),
        "field_coverage_counts": coverage,
        "selection_note": "Hoja de TRABAJO Masari: sólo hojas prospect-like; se excluyeron finanzas, empleados, vacaciones, operaciones y cotizaciones.",
    })

    with (OUT / "prospects.ndjson").open("w", encoding="utf-8", newline="\n") as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False, separators=(",", ":")) + "\n")
    (OUT / "manifest.json").write_text(json.dumps(stats, ensure_ascii=False, indent=2), encoding="utf-8")
    with (OUT / "duplicate-groups.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["dedupe_key", "record_count", "record_keys"])
        for key, values in sorted(duplicate_groups.items()):
            writer.writerow([key, len(values), " | ".join(values)])
    print(json.dumps({
        "records_total": stats["records_total"],
        "records_by_file": stats["records_by_file"],
        "duplicate_group_count": stats["duplicate_group_count"],
        "records_with_mojibake_flag": stats["records_with_mojibake_flag"],
        "field_coverage_counts": stats["field_coverage_counts"],
        "output": str(OUT),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
